"""Plain text from PDF and Word files attached to a chat message (pure Python, no OCR)."""
import base64
import io

LIMIT = 40000          # characters kept per file, so one document can't swamp a small model's context


def extract(name, data_url):
    """Return the file's text. Never raises: errors come back as a short note the model can read."""
    try:
        raw = base64.b64decode(data_url.split(",", 1)[1])
        low = name.lower()
        if low.endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(raw)).pages)
        elif low.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(raw))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
            text = "\n".join(parts)
        else:
            return f"(unsupported file type: {name})"
    except Exception as e:           # a broken or encrypted file must not break the chat turn
        return f"(could not read {name}: {e})"
    text = text.strip()
    if not text:
        return f"(no readable text in {name}; scanned or image-only pages need OCR)"
    return text[:LIMIT] + ("\n…(truncated)" if len(text) > LIMIT else "")
