import base64
import io
import unittest

import documents


def data_url(raw, mime):
    return f"data:{mime};base64," + base64.b64encode(raw).decode()


class DocumentTest(unittest.TestCase):
    def test_broken_pdf_gives_a_readable_note(self):
        out = documents.extract("bad.pdf", data_url(b"not a pdf", "application/pdf"))
        self.assertTrue(out.startswith("(could not read bad.pdf"))

    def test_unsupported_type(self):
        self.assertIn("unsupported", documents.extract("a.exe", data_url(b"x", "application/octet-stream")))

    def test_docx_text_and_tables(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx not installed")
        d = docx.Document()
        d.add_paragraph("Invoice due Friday")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Amount", "4200"
        buf = io.BytesIO()
        d.save(buf)
        out = documents.extract("inv.docx", data_url(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        self.assertIn("Invoice due Friday", out)
        self.assertIn("Amount | 4200", out)


if __name__ == "__main__":
    unittest.main()
